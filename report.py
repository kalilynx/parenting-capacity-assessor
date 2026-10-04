"""Report assembly — stitches drafted sections into a markdown document and exports to .docx."""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import List, Dict

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

import config


def build_markdown(sections: List[Dict[str, str]], case_title: str = "") -> str:
    lines = []
    lines.append("# Parenting Capacity Assessment")
    if case_title:
        lines.append(f"\n**Case:** {case_title}")
    lines.append(f"\n**Prepared by:** Dr. Elena Vasquez")
    lines.append(f"**Date:** {datetime.now().strftime('%d %B %Y')}")
    lines.append("\n*CONFIDENTIAL — FOR COURT PROCEEDINGS ONLY*\n")
    lines.append("---\n")

    for s in sections:
        lines.append(f"## {s['title']}\n")
        lines.append(s["body"].strip())
        lines.append("")

    lines.append("---")
    lines.append("\n**End of Report**")
    lines.append(f"\n*Report generated {datetime.now().strftime('%Y-%m-%d %H:%M')}*")
    return "\n".join(lines)


def save_markdown(md: str, filename: str) -> Path:
    out = config.REPORTS_DIR / filename
    out.write_text(md, encoding="utf-8")
    return out


def _setup_styles(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(6)

    title_style = doc.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(16)
    title_style.font.bold = True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_after = Pt(12)

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Times New Roman"
    h1.font.size = Pt(14)
    h1.font.bold = True
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(6)


def _add_paragraphs(doc: Document, text: str) -> None:
    paragraphs = re.split(r"\n\s*\n", text.strip())
    for para_text in paragraphs:
        para = doc.add_paragraph()
        parts = re.split(r"(\*\*.*?\*\*)", para_text)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = para.add_run(part[2:-2])
                run.bold = True
            else:
                clean = part.replace("**", "")
                para.add_run(clean)


def export_docx(sections: List[Dict[str, str]], case_title: str, filename: str) -> Path:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    _setup_styles(doc)

    doc.add_paragraph("PARENTING CAPACITY ASSESSMENT", style="Title")
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_para.add_run("Prepared by: Dr. Elena Vasquez\n").bold = True
    header_para.add_run("Clinical and Forensic Psychologist\n")
    header_para.add_run(f"Date: {datetime.now().strftime('%d %B %Y')}\n")
    header_run = header_para.add_run("\nCONFIDENTIAL — FOR COURT PROCEEDINGS ONLY")
    header_run.bold = True
    header_run.font.color.rgb = RGBColor(150, 0, 0)

    if case_title:
        case_para = doc.add_paragraph()
        case_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        case_para.add_run(f"Case: {case_title}").italic = True

    doc.add_page_break()

    for s in sections:
        doc.add_heading(s["title"], level=1)
        _add_paragraphs(doc, s["body"])

    doc.add_page_break()
    end_para = doc.add_paragraph()
    end_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    end_para.add_run("— End of Report —").italic = True
    end_para.add_run(f"\n\nGenerated {datetime.now().strftime('%d %B %Y %H:%M')}")

    out = config.REPORTS_DIR / filename
    doc.save(str(out))
    return out
